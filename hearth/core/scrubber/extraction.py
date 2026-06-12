#!/usr/bin/env python3
"""
Module: hearth.core.scrubber.extraction
Purpose: Fail-closed text extraction for the DocScrubber (spec v0.3 §3.3).

Every failure mode — missing file, unsupported format, parser exception,
encrypted PDF, empty extraction, missing parser library — raises
ExtractionError. Nothing here returns "" on failure: an unparseable
document must become a refuse finding upstream (§12 P1-B), and the only
way to guarantee that is to make silence impossible here.

Parsers reuse the repo baseline (PyMuPDF for .pdf, python-docx for .docx);
imports are deferred so a missing library degrades to a refusal for that
format instead of breaking every other format.

Safety Critical: False
Financially Critical: True — a scrub that fails OPEN (silently passes a
  document) is the cardinal failure mode (§1.1); never-suppress-error.

BoundState AI 2026
Classification: BoundState Proprietary

Co-Authored-By: Claude <noreply@anthropic.com>
Co-Authored-By: Dr. Joseph Wilde, BoundState AI
"""

from __future__ import annotations

__safety_critical__: bool = False
__financially_critical__: bool = True

import logging
from pathlib import Path

LOGGER = logging.getLogger("hearth.core.scrubber.extraction")

SUPPORTED_SUFFIXES = (".txt", ".md", ".pdf", ".docx")

# Real contracts are kilobytes to a few megabytes; anything past this is
# either the wrong file or an adversarial input. Refusing beats paging the
# machine to death mid-scrub.
MAX_DOCUMENT_BYTES = 16 * 1024 * 1024


class ExtractionError(Exception):
    """Document text could not be extracted; caller must fail CLOSED."""


def extract_text(path: Path) -> str:
    """Extract document text or raise ExtractionError. Never returns empty."""
    if not path.exists():
        raise ExtractionError(f"document not found: {path}")
    if not path.is_file():
        raise ExtractionError(f"not a regular file: {path}")
    size = path.stat().st_size
    if size > MAX_DOCUMENT_BYTES:
        raise ExtractionError(
            f"{path.name} is {size:,} bytes — exceeds the "
            f"{MAX_DOCUMENT_BYTES:,}-byte scrub limit"
        )

    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        text = _extract_plain_text(path)
    elif suffix == ".pdf":
        text = _extract_pdf(path)
    elif suffix == ".docx":
        text = _extract_docx(path)
    else:
        raise ExtractionError(
            f"unsupported document format '{suffix or '(none)'}': "
            f"supported formats are {', '.join(SUPPORTED_SUFFIXES)}"
        )

    if not text.strip():
        raise ExtractionError(f"no extractable text in {path.name}")
    return text


def _extract_plain_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise ExtractionError(f"cannot read {path.name} as UTF-8 text: {exc}") from exc


def _extract_pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF; deferred so .txt scrubs survive a broken install
    except ImportError as exc:
        raise ExtractionError(
            "PyMuPDF is not installed — cannot parse PDF (failing closed)"
        ) from exc
    try:
        with fitz.open(path) as doc:
            if doc.needs_pass:
                raise ExtractionError(f"{path.name} is password-protected")
            if doc.page_count == 0:
                raise ExtractionError(f"{path.name} contains no pages")
            return "\n".join(page.get_text("text") for page in doc)
    except ExtractionError:
        raise
    except Exception as exc:  # PyMuPDF raises library-specific errors; any
        # parser failure is the same contract: unparseable => refuse (§3.3).
        raise ExtractionError(f"PDF parse failed for {path.name}: {exc}") from exc


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx; deferred (see _extract_pdf)
    except ImportError as exc:
        raise ExtractionError(
            "python-docx is not installed — cannot parse DOCX (failing closed)"
        ) from exc
    try:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ExtractionError(f"DOCX parse failed for {path.name}: {exc}") from exc
