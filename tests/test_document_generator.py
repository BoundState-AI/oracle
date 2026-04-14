"""Tests for markdown-to-block parsing in the document generator."""
import tempfile
from pathlib import Path

from docx import Document

from oracle.document_generator import generate_docx, parse_markdown


def test_parse_markdown_preserves_basic_structure():
    md = """# Title

## Section

Paragraph one
continues here.

- bullet one
1. numbered one

```python
print("hi")
```
"""

    blocks = parse_markdown(md)

    assert [(block.kind, block.text) for block in blocks] == [
        ("h1", "Title"),
        ("h2", "Section"),
        ("paragraph", "Paragraph one continues here."),
        ("bullet", "bullet one"),
        ("numbered", "numbered one"),
        ("code", 'print("hi")'),
    ]


def test_parse_markdown_extracts_gfm_table():
    """GFM pipe tables must become structured rows, not collapsed paragraphs.

    This is the regression test for the bug where every table in an Oracle
    review doc was rendered as one long "Normal" paragraph with raw pipes.
    """
    md = """## Stack

| Component | Technology | Version |
|-----------|-----------|---------|
| API | FastAPI | 0.110+ |
| DB  | PostgreSQL | 16 |

After the table.
"""

    blocks = parse_markdown(md)
    kinds = [b.kind for b in blocks]
    assert kinds == ["h2", "table", "paragraph"]

    table = blocks[1]
    assert table.rows == [
        ["Component", "Technology", "Version"],
        ["API", "FastAPI", "0.110+"],
        ["DB", "PostgreSQL", "16"],
    ]
    assert blocks[2].text == "After the table."


def test_parse_markdown_table_with_alignment_colons():
    """Delimiter rows with alignment colons must still be recognized."""
    md = """| Col A | Col B |
|:------|------:|
| left  | right |
"""

    blocks = parse_markdown(md)
    assert len(blocks) == 1
    assert blocks[0].kind == "table"
    assert blocks[0].rows == [["Col A", "Col B"], ["left", "right"]]


def test_parse_markdown_table_immediately_after_paragraph():
    """A table with no blank line before it must not be swallowed by the
    previous paragraph's continuation loop."""
    md = """Intro prose.
| A | B |
|---|---|
| 1 | 2 |
"""

    blocks = parse_markdown(md)
    assert [b.kind for b in blocks] == ["paragraph", "table"]
    assert blocks[0].text == "Intro prose."
    assert blocks[1].rows == [["A", "B"], ["1", "2"]]


def test_generate_docx_emits_real_word_table():
    """End-to-end: a markdown table in the source must produce a real
    docx table (not a Normal paragraph containing pipe characters).

    Uses tempfile.TemporaryDirectory instead of pytest's tmp_path to
    sidestep Windows permission issues with the shared pytest basetemp.
    """
    md = """# Heading

| Component | Version |
|-----------|---------|
| FastAPI   | 0.110+  |
| Postgres  | 16      |
"""

    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "out.docx"
        generate_docx(md, out, title="T")

        doc = Document(str(out))
        assert len(doc.tables) == 1
        t = doc.tables[0]
        assert len(t.rows) == 3
        assert len(t.columns) == 2
        assert t.rows[0].cells[0].text == "Component"
        assert t.rows[1].cells[0].text == "FastAPI"
        assert t.rows[2].cells[1].text == "16"

        # And crucially, no paragraph in the body should still contain a
        # raw pipe-row (which was the bug signature).
        for p in doc.paragraphs:
            assert not (p.text.count("|") >= 3 and "---" in p.text), (
                f"raw pipe-table leaked into paragraph: {p.text!r}"
            )
